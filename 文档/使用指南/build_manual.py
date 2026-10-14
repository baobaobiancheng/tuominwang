#!/usr/bin/env python3
"""Build a screenshot-based manual from manual.json without adding factual content.

Run with the bundled Python runtime. Image paths are relative to the JSON file.
Schema: {title, intro: str|list[str], updated?: str, chapters: [
  {id?: str, title, intro?: str|list[str], steps?: list[str], result?: str,
   notes?: list[str], images?: [{path, caption}], image?: str}]}
"""
from __future__ import annotations
import argparse
import base64
import hashlib
import html
import json
import mimetypes
import os
import subprocess
import sys
from pathlib import Path
from urllib.parse import quote

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from PIL import Image


def paragraphs(value):
    return [value] if isinstance(value, str) and value else value or []


def chapter_images(chapter, source):
    images = chapter.get("images", [])
    if chapter.get("image"):
        images = [chapter["image"], *images]
    for item in images:
        item = {"path": item, "caption": ""} if isinstance(item, str) else item
        path = (source.parent / item["path"]).resolve()
        if not path.is_file():
            raise FileNotFoundError(path)
        with Image.open(path) as im:
            width, height = im.size
            im.verify()
        yield {"path": path, "relative": item["path"], "caption": item.get("caption", ""), "width": width, "height": height}


def set_font(style, size, bold=False):
    style.font.name = "Arial"
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for key in ("eastAsia", "ascii", "hAnsi"):
        fonts.set(qn(f"w:{key}"), "Heiti SC" if key == "eastAsia" else "Arial")
    for key in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        fonts.attrib.pop(qn(f"w:{key}"), None)


def add_bookmark(paragraph, name, bookmark_id):
    start, end = OxmlElement("w:bookmarkStart"), OxmlElement("w:bookmarkEnd")
    start.set(qn("w:id"), str(bookmark_id)); start.set(qn("w:name"), name)
    end.set(qn("w:id"), str(bookmark_id))
    paragraph._p.insert(0, start); paragraph._p.append(end)


def add_link(paragraph, text, anchor):
    link = OxmlElement("w:hyperlink"); link.set(qn("w:anchor"), anchor)
    run, value = OxmlElement("w:r"), OxmlElement("w:t")
    value.text = text; run.append(value); link.append(run); paragraph._p.append(link)


def build(source: Path, output: Path):
    data = json.loads(source.read_text(encoding="utf-8"))
    if not data.get("title") or not data.get("chapters"):
        raise ValueError("manual.json needs title and nonempty chapters")
    chapters = []
    for index, chapter in enumerate(data["chapters"], 1):
        if not chapter.get("title"):
            raise ValueError(f"Chapter {index} has no title")
        chapters.append({**chapter, "anchor": f"chapter-{index}", "number": index, "pictures": list(chapter_images(chapter, source))})
    output.mkdir(parents=True, exist_ok=True)
    stem = data.get("output_stem", "脱敏汪使用指南")
    doc = Document()
    section = doc.sections[0]
    landscape = data.get("orientation", "landscape") == "landscape"
    section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    section.page_width, section.page_height = (Mm(297), Mm(210)) if landscape else (Mm(210), Mm(297))
    section.top_margin = section.bottom_margin = Mm(16)
    section.left_margin = section.right_margin = Mm(16)
    section.footer_distance = Mm(7)
    for name, size, bold in [("Normal", 10.5, False), ("Title", 24, True), ("Heading 1", 17, True), ("Heading 2", 12, True), ("Caption", 9, False), ("List Number", 10.5, False)]:
        set_font(doc.styles[name], size, bold)
        pf = doc.styles[name].paragraph_format
        pf.line_spacing = 1.25; pf.space_after = Pt(6)
        for tag in ("w:pBdr", "w:shd"):
            for element in doc.styles[name].element.findall(f".//{qn(tag)}"):
                element.getparent().remove(element)
    doc.styles["Heading 1"].paragraph_format.space_after = Pt(9)
    doc.styles["Normal"].paragraph_format.widow_control = True
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement("w:fldSimple"); field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    doc.core_properties.title = data["title"]
    doc.core_properties.subject = ""
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.add_paragraph(data["title"], "Title")
    for text in paragraphs(data.get("intro")): doc.add_paragraph(text)
    if data.get("updated"): doc.add_paragraph(data["updated"])
    doc.add_heading("目录", 1)
    for chapter in chapters:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.paragraph_format.line_spacing = 1.05
        add_link(paragraph, f'{chapter["number"]} {chapter["title"]}', chapter["anchor"].replace("-", "_"))
    escaped = html.escape
    html_parts = [f'<!doctype html><html lang="zh-CN"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>{escaped(data["title"])}</title>', '''<style>
*{box-sizing:border-box}body{margin:0;color:#111;background:#fff;font:16px/1.8 -apple-system,BlinkMacSystemFont,"PingFang SC","Microsoft YaHei",sans-serif}main{max-width:1100px;margin:0 auto;padding:40px 32px 90px}h1,h2,h3{color:#000;line-height:1.4}h1{font-size:32px;margin:0 0 24px}h2{font-size:24px;margin:0 0 20px}p,ol,ul{margin:12px 0 20px}li{margin:8px 0}nav{margin:32px 0 48px}nav ol{columns:2;column-gap:48px}nav li{break-inside:avoid}a{color:inherit;text-underline-offset:4px}section{margin:64px 0;scroll-margin-top:24px}figure{margin:28px 0 40px}img{display:block;max-width:100%;height:auto;margin:auto;border:1px solid #ddd;cursor:zoom-in}img.zoomed{max-width:none;cursor:zoom-out}figure:has(img.zoomed){overflow:auto}figcaption{font-size:14px;margin:12px 0;color:#333}.back{font-size:14px}.result{margin-top:18px}footer{font-size:13px;color:#444;margin-top:60px}@media(max-width:700px){main{padding:24px 18px}nav ol{columns:1}h1{font-size:26px}h2{font-size:21px}}@media print{main{padding:0;max-width:none}.back{display:none}section{break-before:page;margin:0}figure{break-inside:avoid}img{max-height:210mm;max-width:100%!important;object-fit:contain;border:0}nav ol{columns:1}}
</style></head><body><main id="top">''', f'<h1>{escaped(data["title"])}</h1>']
    markdown = [f'# {data["title"]}', ""]
    for text in paragraphs(data.get("intro")):
        html_parts.append(f'<p>{escaped(text)}</p>'); markdown.extend([text, ""])
    if data.get("updated"):
        html_parts.append(f'<p>{escaped(data["updated"])}</p>'); markdown.extend([data["updated"], ""])
    html_parts.append('<nav aria-label="目录"><h2>目录</h2><ol>')
    for chapter in chapters: html_parts.append(f'<li><a href="#{chapter["anchor"]}">{escaped(chapter["title"])}</a></li>')
    html_parts.append('</ol></nav>')
    inventory = []
    for chapter in chapters:
        doc.add_page_break()
        heading = doc.add_heading(f'{chapter["number"]} {chapter["title"]}', 1)
        add_bookmark(heading, chapter["anchor"].replace("-", "_"), chapter["number"])
        html_parts.append(f'<section id="{chapter["anchor"]}"><h2>{chapter["number"]} {escaped(chapter["title"])}</h2>')
        markdown.extend([f'## {chapter["number"]} {chapter["title"]}', ""])
        for text in paragraphs(chapter.get("intro")):
            doc.add_paragraph(text); html_parts.append(f'<p>{escaped(text)}</p>'); markdown.extend([text, ""])
        if chapter.get("steps"):
            html_parts.append('<ol>')
            for step, text in enumerate(chapter["steps"], 1):
                doc.add_paragraph(f'{step}. {text}')
                html_parts.append(f'<li>{escaped(text)}</li>'); markdown.append(f'{step}. {text}')
            html_parts.append('</ol>'); markdown.append("")
        if chapter.get("result"):
            p = doc.add_paragraph(); p.add_run("操作结果  ").bold = True; p.add_run(chapter["result"])
            html_parts.append(f'<p class="result"><strong>操作结果</strong> {escaped(chapter["result"])}</p>')
            markdown.extend([f'操作结果：{chapter["result"]}', ""])
        for text in chapter.get("notes", []):
            doc.add_paragraph(text); html_parts.append(f'<p>{escaped(text)}</p>'); markdown.extend([text, ""])
        for attachment in chapter.get("files", []):
            relative = attachment["path"]
            file_path = (source.parent / relative).resolve()
            raw = file_path.read_bytes()
            label = attachment.get("label", file_path.name)
            paragraph = doc.add_paragraph("实际文件：")
            relationship = paragraph.part.relate_to(quote(relative), "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
            link = OxmlElement("w:hyperlink"); link.set(qn("r:id"), relationship)
            run, value = OxmlElement("w:r"), OxmlElement("w:t")
            value.text = label; run.append(value); link.append(run); paragraph._p.append(link)
            mime = mimetypes.guess_type(file_path.name)[0] or "application/octet-stream"
            uri = f'data:{mime};base64,' + base64.b64encode(raw).decode("ascii")
            html_parts.append(f'<p>实际文件：<a href="{uri}" download="{escaped(file_path.name, quote=True)}">{escaped(label)}</a></p>')
            markdown.extend([f'实际文件：[{label}]({quote(relative)})', ""])
        for image_index, picture in enumerate(chapter["pictures"]):
            if image_index:
                doc.add_page_break()
                doc.add_heading(f'{chapter["number"]} {chapter["title"]} 操作界面', 2)
            # Preserve every source pixel and its aspect ratio. No crop or screenshot redraw.
            maximum_height = (151 if landscape else 215) if image_index else (151 if landscape else 183)
            width_mm = min(265 if landscape else 178, maximum_height * picture["width"] / picture["height"])
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = bool(picture["caption"])
            shape = p.add_run().add_picture(str(picture["path"]), width=Mm(width_mm))
            shape._inline.docPr.set("descr", picture["caption"] or chapter["title"])
            if picture["caption"]:
                caption = doc.add_paragraph(picture["caption"], "Caption"); caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            raw = picture["path"].read_bytes()
            mime = mimetypes.guess_type(picture["path"].name)[0] or "image/png"
            uri = f'data:{mime};base64,' + base64.b64encode(raw).decode("ascii")
            html_parts.append(f'<figure><img src="{uri}" alt="{escaped(picture["caption"] or chapter["title"], quote=True)}" width="{picture["width"]}" height="{picture["height"]}" onclick="this.classList.toggle(\'zoomed\')"><figcaption>{escaped(picture["caption"])}</figcaption></figure>')
            markdown.extend([f'![{picture["caption"] or chapter["title"]}]({picture["relative"]})', ""])
            inventory.append({"chapter": chapter["number"], "path": str(picture["path"]), "sha256": hashlib.sha256(raw).hexdigest(), "width": picture["width"], "height": picture["height"]})
        html_parts.append('<p class="back"><a href="#top">返回目录</a></p></section>')
    html_parts.append('</main></body></html>')
    outputs = {"docx": output / f'{stem}.docx', "html": output / f'{stem}.html', "markdown": output / f'{stem}.md'}
    doc.save(outputs["docx"])
    outputs["html"].write_text("\n".join(html_parts), encoding="utf-8")
    outputs["markdown"].write_text("\n".join(markdown), encoding="utf-8")
    result = {"source": str(source), "chapters": len(chapters), "images": inventory, "outputs": {key: str(path) for key, path in outputs.items()}}
    (output / "build-result.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"chapters": len(chapters), "images": len(inventory), "outputs": result["outputs"]}, ensure_ascii=False))


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", nargs="?", type=Path, default=Path(__file__).with_name("manual.json"))
    parser.add_argument("--output-dir", type=Path)
    parser.add_argument("--render", action="store_true", help="Render with bundled LibreOffice and fontconfig")
    parser.add_argument("--render-dir", type=Path)
    args = parser.parse_args()
    source = args.source.resolve()
    output = (args.output_dir or source.parent).resolve()
    build(source, output)
    if args.render:
        bundled = Path.home() / ".cache/codex-runtimes/codex-primary-runtime/dependencies"
        python = bundled / "python/bin/python3"
        skill = Path.home() / ".codex/plugins/cache/openai-primary-runtime/documents/26.1011.11113/skills/documents"
        fontconfig = bundled / "native/libreoffice-headless/libreoffice/LibreOfficeDev.app/Contents/Resources/fontconfig/fonts.conf"
        for required in (python, skill / "render_docx.py", fontconfig):
            if not required.is_file(): raise FileNotFoundError(required)
        env = {**os.environ, "FONTCONFIG_FILE": str(fontconfig)}
        stem = json.loads(source.read_text(encoding="utf-8")).get("output_stem", "脱敏汪使用指南")
        subprocess.run([str(python), str(skill / "render_docx.py"), str(output / f"{stem}.docx"), "--output_dir", str((args.render_dir or output / "排版核验").resolve()), "--emit_pdf"], env=env, check=True)
